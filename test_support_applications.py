"""Isolation, persistence, jobs, versioning and real document structure regressions."""
import copy
from contextlib import closing
import hashlib
import http.client
import io
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
import urllib.error
from unittest.mock import patch
import zipfile

from docx import Document
from lxml import etree
import openpyxl

import app
import support_applications as support
import support_documents as docs


def docx_form():
    doc = Document(); doc.add_heading('지원사업 신청서', 0)
    table = doc.add_table(rows=3, cols=2); table.style = 'Table Grid'
    for row, label in zip(table.rows, ['기업명', '대표자', '추진계획']): row.cells[0].text = label
    doc.add_paragraph('신청인:                    (서명 또는 인)')
    out = io.BytesIO(); doc.save(out); return out.getvalue()


class SupportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.env = patch.dict(os.environ, {'SUPPORT_DB_PATH': str(Path(self.temp.name)/'support.db'),
            'SUPPORT_ACCESS_KEY': 'test-key-only', 'OPENAI_API_KEY': 'test-no-network', 'SUPPORT_WORKER_ENABLED': 'false'})
        self.env.start(); self.addCleanup(self.env.stop); support.init_db()
        self.case = {'id': 'bizinfo:PBLN_123', 'title': '검증용 지원사업', 'url': 'https://www.bizinfo.go.kr/sii/siia/selectSIIA200Detail.do?pblancId=PBLN_123', 'application_period': '2026.09.01 ~ 2026.09.30'}
        with closing(support.connect()) as db, db: db.execute('INSERT INTO support_cases VALUES(?,?,\'ready\',\'\',?)', (self.case['id'],support.dumps(self.case),support.now()))

    def configured(self):
        profile = support.profile(); profile['data']['company']['representative']='검증대표'
        profile['data']['plans']=[{'id':'plan','name':'기준계획','overview':'검증용 업무 기록 자동화'}]
        support.save_profile(profile)
        asset = support.store_asset(self.case['id'],'신청서.docx',docx_form(),kind='form')[0]
        return asset

    def generate(self):
        asset = self.configured()
        created = support.create_draft({'case_id':self.case['id'],'branch_id':'incheon','plan_id':'plan','asset_ids':[asset['id']]})
        analysis = json.loads(support.get_asset(asset['id'])['analysis'])
        target = next(t for t in analysis['targets'] if '기업명' in t['context'] and not t['text'])
        result = {'fields':[{'target_id':target['id'],'label':'기업명','value':'주식회사 콤파스원','kind':'fact','sources':['company.name'],'note':''}], 'warnings':['제출 전 확인']}
        with closing(support.connect()) as db, db: job = db.execute("SELECT * FROM support_jobs WHERE kind='draft'").fetchone()
        with patch.object(support,'ai_json',return_value=result),patch.object(support,'collect_case'),patch.object(support,'review_fields'): support.run_job(job)
        return support.get_draft(created['id']), asset

    def test_profile_revision_conflict_and_company_branch_separation(self):
        p = support.profile(); p['data']['company']['address']='성남 본사 주소'; p['data']['branches'][3]['address']='인천 사업장 주소'
        saved = support.save_profile(p); self.assertEqual(saved['revision'],2)
        with self.assertRaises(support.wiki_chat.ChatError): support.save_profile(p)
        self.assertNotEqual(saved['data']['company']['address'],saved['data']['branches'][3]['address'])
        support.init_db(); self.assertEqual(support.profile()['revision'],2)

    def test_document_filling_preserves_labels_tables_and_signature(self):
        raw = docx_form(); result = docs.inspect_document('신청서.docx',raw)
        target = next(t for t in result['targets'] if '대표자' in t['context'] and not t['text'])
        suffix, filled = docs.fill_document('신청서.docx',raw,[{'target_id':target['id'],'value':'검증대표'}])
        self.assertEqual(suffix,'.docx'); doc = Document(io.BytesIO(filled))
        self.assertEqual(doc.tables[0].cell(1,0).text,'대표자'); self.assertEqual(doc.tables[0].cell(1,1).text,'검증대표')
        self.assertIn('(서명 또는 인)',doc.paragraphs[-1].text)
        self.assertEqual(len(doc.tables[0].rows),3)
        label = next(t for t in result['targets'] if t['text'] == '대표자')
        self.assertFalse(label['editable']); self.assertEqual(target['input_label'], '대표자')
        with self.assertRaises(ValueError): docs.fill_document('신청서.docx', raw, [{'target_id':label['id'],'value':'이름'}])

    def test_independent_review_removes_unsupported_claim_and_requires_every_field(self):
        fields = [{'id':'f1','target_id':'t1','label':'공급기업 역량','value':'이미 구축 경험이 풍부합니다.','kind':'narrative','sources':['plan.overview'],'note':''}]
        result = {'reviews':[{'id':'f1','supported':False,'reason':'공급기업의 실제 수행경험을 등록해주세요.'}]}
        with patch.object(support, 'ai_json', return_value=result): support.review_fields(fields, {'plan.overview':'도입 계획'}, {})
        self.assertEqual(fields[0]['value'],'');self.assertEqual(fields[0]['kind'],'missing');self.assertIn('공급기업',fields[0]['note'])
        fields[0]['value']='다시 검토할 내용'
        with patch.object(support, 'ai_json', return_value={'reviews':[]}):
            with self.assertRaises(ValueError): support.review_fields(fields, {}, {})

    def test_identifiers_use_adjacent_labels_and_supplier_is_not_applicant(self):
        targets=[{'id':'company','input_label':'기업명','left_labels':['기업명'],'above':''},
                 {'id':'rep','input_label':'성명','left_labels':['기업명','대표자','성명'],'above':''},
                 {'id':'contact','input_label':'성명','left_labels':['사업자등록번호','실 무 자','성명'],'above':''},
                 {'id':'supplier','input_label':'기업명','left_labels':['AX 공급기업','기업명'],'above':''}]
        fields=support.exact_fact_fields(targets,{'company.name':'회사','company.representative':'대표이름'})
        self.assertEqual({f['target_id']:f['value'] for f in fields},{'company':'회사','rep':'대표이름','contact':''})

    def test_rewrite_uses_grounding_review_and_preserves_previous_version(self):
        draft,_=self.generate();field=draft['data']['documents'][0]['fields'][0]
        result={k:field[k] for k in support.FIELD_SCHEMA['required']};result['value']='근거 없는 회사명'
        review={'reviews':[{'id':field['id'],'supported':False,'reason':'등록 회사명을 확인해주세요.'}]}
        with patch.object(support,'ai_json',side_effect=[result,review]):
            saved=support.rewrite_field({'id':draft['id'],'field_id':field['id'],'instruction':'다시 작성'})
        self.assertEqual(saved['version'],2);self.assertEqual(saved['data']['documents'][0]['fields'][0]['value'],'')
        self.assertEqual(support.get_draft(draft['id'])['data']['documents'][0]['fields'][0]['value'],'주식회사 콤파스원')

    def test_interest_collect_does_not_call_ai_and_retries_deduplicate(self):
        with patch.object(support.agency_news,'report',return_value={'items':[self.case|{'kind':'support','preference':'interested'}]}), patch.object(support,'ai_json') as model:
            self.assertTrue(support.ensure_case(self.case['id'],True)); self.assertFalse(support.ensure_case(self.case['id'],True)); model.assert_not_called()
        first=support.store_asset(self.case['id'],'신청서.docx',docx_form())
        raw=support.get_asset(first[0]['id'])['content']
        second=support.store_asset(self.case['id'],'신청서.docx',raw)
        self.assertEqual(first[0]['id'],second[0]['id'])

    def test_jobs_are_explicit_and_repeat_generation_is_rejected(self):
        asset=self.configured(); body={'case_id':self.case['id'],'branch_id':'incheon','plan_id':'plan','asset_ids':[asset['id']]}
        with patch.object(support,'ai_json') as model:
            support.create_draft(body);model.assert_not_called()
            with self.assertRaises(support.wiki_chat.ChatError):support.create_draft(body)

    def test_generation_export_versions_and_profile_snapshots(self):
        draft,asset=self.generate();self.assertEqual(draft['status'],'ready',draft['error'])
        field=draft['data']['documents'][0]['fields'][0]
        saved=support.save_draft({'id':draft['id'],'values':{field['id']:'사용자 수정 회사명'}})
        self.assertEqual(saved['version'],2);self.assertEqual(support.get_draft(draft['id'])['data']['documents'][0]['fields'][0]['value'],'주식회사 콤파스원')
        profile=support.profile();profile['data']['company']['name']='변경된 회사명';support.save_profile(profile)
        self.assertEqual(support.get_draft(saved['id'])['snapshot']['company']['name'],'주식회사 콤파스원')
        name,raw=support.draft_export(saved['id'],asset['id']);self.assertTrue(name.endswith('.docx'))
        self.assertEqual(Document(io.BytesIO(raw)).tables[0].cell(0,1).text,'사용자 수정 회사명')
        name,raw=support.draft_export(saved['id'])
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            self.assertEqual(len(archive.namelist()),3);self.assertIn('보완사항.txt',archive.namelist())

    def test_unsupported_ai_target_does_not_become_completed(self):
        asset=self.configured();support.create_draft({'case_id':self.case['id'],'branch_id':'incheon','plan_id':'plan','asset_ids':[asset['id']]})
        with closing(support.connect()) as db, db:job=db.execute('SELECT * FROM support_jobs').fetchone()
        with patch.object(support,'ai_json',return_value={'fields':[{'target_id':'invented','label':'대표자','value':'없음','kind':'fact','sources':[],'note':''}],'warnings':[]}),patch.object(support,'collect_case'):support.run_job(job)
        with closing(support.connect()) as db, db:self.assertEqual(db.execute('SELECT status FROM support_drafts').fetchone()[0],'failed')
        self.assertEqual(len(support.assets(self.case['id'])),1)

    def test_zip_bombs_paths_invalid_magic_and_formula_preservation(self):
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as archive:archive.writestr('../escape.docx',b'no')
        with self.assertRaises(ValueError):docs.validate_file('bad.zip',out.getvalue())
        with self.assertRaises(ValueError):docs.validate_file('bad.hwp',b'<html>login</html>')
        book=openpyxl.Workbook();book.active['A1']='금액';book.active['B1']='=SUM(C1:C2)';book.active['C1']=100
        out=io.BytesIO();book.save(out);raw=out.getvalue()
        self.assertNotIn('Sheet!B1',{t['id'] for t in docs.inspect_document('예산.xlsx',raw)['targets']})
        _,filled=docs.fill_document('예산.xlsx',raw,[{'target_id':'Sheet!C1','value':'200'}]);self.assertEqual(openpyxl.load_workbook(io.BytesIO(filled)).active['B1'].value,'=SUM(C1:C2)')
        with self.assertRaises(ValueError):docs.fill_document('예산.xlsx',raw,[{'target_id':'Sheet!B1','value':'3'}])

    def test_private_dns_and_unsafe_destination_rejected(self):
        for url in ['http://www.bizinfo.go.kr/', 'https://user:pass@example.com/', 'file:///etc/passwd']:
            with self.assertRaises(ValueError):docs.public_get(url)
        with patch.object(docs.socket,'getaddrinfo',return_value=[(2,1,6,'',('127.0.0.1',443))]):
            with self.assertRaises(ValueError):docs.PublicHTTPS('example.com').connect()

    def test_auth_download_head_and_cross_origin_write(self):
        asset=self.configured();server=app.ThreadingHTTPServer(('127.0.0.1',0),app.App)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        self.addCleanup(server.server_close);self.addCleanup(server.shutdown)
        def request(method,route,body=None,cookie=None,origin=None):
            c=http.client.HTTPConnection('127.0.0.1',server.server_port);headers={}
            if cookie:headers['Cookie']=cookie
            if origin:headers['Origin']=origin
            if body is not None:headers['Content-Type']='application/json'
            c.request(method,'/api/support/'+route,None if body is None else json.dumps(body),headers);r=c.getresponse();raw=r.read();result=(r.status,dict(r.getheaders()),raw);c.close();return result
        self.assertEqual(request('GET','profile')[0],401)
        self.assertEqual(request('HEAD','assets/file?id='+asset['id'])[0],401)
        code,headers,_=request('POST','login',{'key':'test-key-only'});self.assertEqual(code,200)
        cookie=headers['Set-Cookie'];self.assertIn('HttpOnly',cookie);self.assertIn('SameSite=Strict',cookie)
        self.assertEqual(request('GET','profile',cookie=cookie)[0],200)
        self.assertEqual(request('POST','profile',support.profile(),cookie=cookie,origin='https://evil.example')[0],403)
        code,headers,raw=request('GET','assets/file?id='+asset['id'],cookie=cookie);self.assertEqual(code,200);self.assertEqual(headers['Cache-Control'],'no-store');self.assertTrue(raw.startswith(b'PK'))
        self.assertEqual(request('POST','logout',{},cookie=cookie)[0],200);self.assertEqual(request('GET','profile',cookie=cookie)[0],401)


class AITransportTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ,{'OPENAI_API_KEY':'test-no-network','SUPPORT_OPENAI_MODEL':'gpt-4.1-mini'})
        self.env.start();self.addCleanup(self.env.stop)
        self.delay=patch.object(support.time,'sleep');self.delay.start();self.addCleanup(self.delay.stop)
        self.schema={'type':'object','properties':{'ok':{'type':'boolean'}},'required':['ok'],'additionalProperties':False}

    def response(self, status='completed', reason=None, text='{"ok":true}', refusal=False):
        result={'id':'response-test','status':status,'incomplete_details':{'reason':reason} if reason else None,
                'usage':{'input_tokens':42,'output_tokens':12000 if reason=='max_output_tokens' else 8},
                'output':[{'type':'message','content':[{'type':'refusal','refusal':'refused'} if refusal else {'type':'output_text','text':text}]}]}
        return io.BytesIO(json.dumps(result).encode())

    def test_truncated_response_retries_with_larger_budget_without_exposing_partial_json(self):
        with patch.object(support.urllib.request,'urlopen',side_effect=[self.response('incomplete','max_output_tokens','{"ok":'),self.response()]) as send, self.assertLogs(support.LOG,level='INFO') as logs:
            result=support.ai_json('instructions',{'private':'company-data-must-not-be-logged'},self.schema,'support_application')
        self.assertEqual(result,{'ok':True});self.assertEqual(send.call_count,2)
        payloads=[json.loads(c.args[0].data) for c in send.call_args_list]
        self.assertEqual([p['max_output_tokens'] for p in payloads],[12000,24000])
        self.assertTrue(all(p['store'] is False for p in payloads))
        self.assertIn('max_output_tokens',' '.join(logs.output));self.assertNotIn('company-data-must-not-be-logged',' '.join(logs.output))

    def test_retry_is_bounded_and_reports_incomplete_separately(self):
        with patch.object(support.urllib.request,'urlopen',side_effect=[self.response('incomplete','max_output_tokens'),self.response('incomplete','max_output_tokens')]) as send:
            with self.assertRaisesRegex(ValueError,'자동 재시도'):support.ai_json('i',{},self.schema,'test')
        self.assertEqual(send.call_count,2)

    def test_filtered_or_refused_content_is_not_automatically_retried(self):
        for response in [self.response('incomplete','content_filter'),self.response(refusal=True)]:
            with self.subTest(),patch.object(support.urllib.request,'urlopen',return_value=response) as send:
                with self.assertRaisesRegex(ValueError,'자료와 작성 요청'):support.ai_json('i',{},self.schema,'test')
                self.assertEqual(send.call_count,1)

    def test_transport_and_invalid_json_recover_once(self):
        for failure in [TimeoutError(),http.client.IncompleteRead(b'partial'),self.response(text='{"ok":')]:
            with self.subTest(failure=type(failure).__name__),patch.object(support.urllib.request,'urlopen',side_effect=[failure,self.response()]) as send:
                self.assertEqual(support.ai_json('i',{},self.schema,'test'),{'ok':True});self.assertEqual(send.call_count,2)

    def test_server_errors_retry_but_auth_and_quota_errors_do_not(self):
        for code,expected in [(503,2),(401,1),(429,1)]:
            error=urllib.error.HTTPError('https://api.openai.com/v1/responses',code,'API error',{},None)
            with self.subTest(code=code),patch.object(support.urllib.request,'urlopen',side_effect=[error,self.response()]) as send:
                if expected==2:self.assertEqual(support.ai_json('i',{},self.schema,'test'),{'ok':True})
                else:
                    with self.assertRaises(ValueError):support.ai_json('i',{},self.schema,'test')
                self.assertEqual(send.call_count,expected)


if __name__=='__main__':unittest.main()
