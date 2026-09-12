"""Bounded public attachment retrieval and structure-preserving application documents."""
from __future__ import annotations

import copy
import hashlib
import http.client
import io
import ipaddress
import json
from pathlib import Path
import re
import socket
import ssl
import subprocess
import sys
import tempfile
import urllib.parse
import zipfile

from lxml import etree, html

MAX_FILE = 12 * 1024 * 1024
MAX_EXPANDED = 60 * 1024 * 1024
FORMATS = {'.hwp', '.hwpx', '.docx', '.xlsx', '.pdf', '.txt', '.md', '.zip'}
HP = 'http://www.hancom.co.kr/hwpml/2011/paragraph'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def clean_name(value):
    value = re.sub(r'[\x00-\x1f<>:"/\\|?*]', '_', str(value))
    return value.strip(' .')[:180] or '첨부파일'


class PublicHTTPS(http.client.HTTPSConnection):
    def connect(self):
        addresses = socket.getaddrinfo(self.host, self.port, type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
            raise ValueError('공개 기관의 인터넷 주소만 읽을 수 있습니다.')
        # Pin the checked address, keeping certificate validation against the hostname.
        self.sock = socket.create_connection((addresses[0][4][0], self.port), self.timeout)
        self.sock = self._context.wrap_socket(self.sock, server_hostname=self.host)


def public_get(url, limit=MAX_FILE):
    for _ in range(5):
        parts = urllib.parse.urlsplit(url)
        if parts.scheme != 'https' or not parts.hostname or parts.username or parts.password or parts.port not in (None, 443, 8443):
            raise ValueError('HTTPS 공식 공고 주소를 확인해주세요.')
        conn = PublicHTTPS(parts.hostname, parts.port or 443, timeout=25, context=ssl.create_default_context())
        try:
            target = urllib.parse.quote(parts.path or '/', safe='/%:@') + ('?' + parts.query if parts.query else '')
            conn.request('GET', target, headers={'User-Agent': 'TheVidaSupportPreparation/1.0', 'Accept-Encoding': 'identity'})
            response = conn.getresponse()
            if response.status in (301, 302, 303, 307, 308):
                url = urllib.parse.urljoin(url, response.getheader('Location', ''))
                continue
            if response.status != 200:
                raise ValueError(f'공고 서버에서 파일을 받지 못했습니다. (HTTP {response.status})')
            raw = response.read(limit + 1)
            if len(raw) > limit:
                raise ValueError('파일이 허용 크기를 초과합니다. 12MB 이하 파일을 직접 등록해주세요.')
            return raw, dict(response.getheaders()), url
        finally:
            conn.close()
    raise ValueError('공고 주소의 이동이 너무 많습니다. 원문에서 파일을 확인해주세요.')


def parse_attachments(raw, url):
    # Bizinfo declares UTF-8 in its HTML; lxml otherwise assumes Latin-1 for bytes.
    doc = html.fromstring(raw.decode('utf-8', errors='replace'))
    found, seen = [], set()
    for a in doc.xpath('//a[@href]'):
        href = a.get('href', '')
        label = a.get('title', '') or a.text_content()
        if not ('fileDown.do' in href or re.search(r'\.(hwp|hwpx|docx|xlsx|pdf|zip)(?:[?#]|$)', href, re.I)
                or ('다운로드' in label and re.search(r'download|fileDown|downFile', href, re.I))):
            continue
        target = urllib.parse.urljoin(url, href)
        if not target.startswith('https://') or target in seen:
            continue
        seen.add(target)
        name = re.sub(r'^첨부파일\s*|\s*다운로드$', '', label).strip()
        if not Path(name).suffix.lower() in FORMATS:
            name = a.getparent().text_content().strip()
            match = re.search(r'([^\n]+\.(?:hwpx|hwp|docx|xlsx|pdf|zip))', name, re.I)
            name = match[1].strip() if match else urllib.parse.unquote(Path(urllib.parse.urlsplit(target).path).name)
        found.append({'name': clean_name(name), 'url': target})
    source = doc.xpath('//a[@id="barogagi"]/@href')
    return found[:24], (urllib.parse.urljoin(url, source[0]) if source else None)


def category(name):
    if re.search('동의|확약|서약', name):
        return 'consent' if not re.search('신청|계획', name) else 'form'
    if re.search('신청|계획|양식|서식|참가|제안|지원서', name):
        return 'form'
    if re.search('공고|포스터|안내|요령', name):
        return 'notice'
    return 'reference'


def safe_zip(raw):
    archive = zipfile.ZipFile(io.BytesIO(raw))
    members = archive.infolist()
    if len(members) > 1500 or sum(i.file_size for i in members) > MAX_EXPANDED:
        raise ValueError('압축파일의 항목 수 또는 압축 해제 크기가 너무 큽니다.')
    for item in members:
        p = Path(item.filename.replace('\\', '/'))
        if p.is_absolute() or '..' in p.parts or item.flag_bits & 1:
            raise ValueError('암호화되었거나 올바르지 않은 압축파일입니다.')
    return archive


def validate_file(name, raw):
    suffix = Path(name).suffix.lower()
    if suffix not in FORMATS or not raw or len(raw) > MAX_FILE:
        raise ValueError('HWP, HWPX, DOCX, XLSX, PDF, TXT, MD, ZIP 파일을 12MB 이하로 등록해주세요.')
    if suffix == '.hwp' and not raw.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise ValueError('한글 HWP 파일의 내용을 확인해주세요.')
    if suffix == '.pdf' and not raw.startswith(b'%PDF-'):
        raise ValueError('PDF 파일의 내용을 확인해주세요.')
    if suffix in {'.hwpx', '.docx', '.xlsx', '.zip'}:
        with safe_zip(raw) as archive:
            required = {'.hwpx': 'Contents/section0.xml', '.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml'}
            if suffix in required and required[suffix] not in archive.namelist():
                raise ValueError('확장자에 맞는 문서 구조를 찾지 못했습니다.')
    return suffix


def convert_hwp(raw):
    with tempfile.TemporaryDirectory(prefix='support-hwp-') as temp:
        src, dest = Path(temp) / 'source.hwp', Path(temp) / 'output.hwpx'
        src.write_bytes(raw)
        # Isolate malformed binary parsers, bound conversion time, do not run document macros.
        result = subprocess.run([sys.executable, '-c',
            'from pyhwpxlib.hwp2hwpx import convert; import sys; convert(sys.argv[1],sys.argv[2])',
            str(src), str(dest)], capture_output=True, timeout=45)
        if result.returncode or not dest.exists():
            raise ValueError('HWP 변환을 완료하지 못했습니다. 한글에서 HWPX로 저장해 등록해주세요.')
        converted = dest.read_bytes()
        validate_file('converted.hwpx', converted)
        return converted


def xml_root(raw):
    return etree.fromstring(raw, parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def local_name(node):
    return etree.QName(node).localname


def editable_text(text):
    if re.search('[□☐☑■✓]|서명|날인|동의합니다|동의함', text): return False
    return not text.strip() or bool(re.search(r'※|예\s*[:：]|작성|기재|입력|_{2,}|\{\{', text))


def node_text(node, ns):
    def text_content(t):
        parts = [t.text or '']
        for child in t:
            if local_name(child) in ('lineBreak', 'br'): parts.append('\n')
            elif local_name(child) == 'tab': parts.append('\t')
            parts.append(child.tail or '')
        return ''.join(parts)
    return '\n'.join(text_content(t) for t in node.findall('.//{%s}t' % ns)).strip()


def xml_targets(raw, suffix):
    ns = HP if suffix == '.hwpx' else W
    targets, nodes, roots = [], {}, {}
    with safe_zip(raw) as archive:
        paths = sorted(n for n in archive.namelist() if re.fullmatch(r'Contents/section\d+\.xml', n)) if suffix == '.hwpx' else ['word/document.xml']
        for path in paths:
            root = xml_root(archive.read(path)); roots[path] = root
            table_no = 0
            for table in root.findall('.//{%s}tbl' % ns):
                rows = table.findall('{%s}tr' % ns)
                layout = []
                for r, row in enumerate(rows):
                    column = 0
                    for cell in row.findall('{%s}tc' % ns):
                        if suffix == '.hwpx':
                            addr, span = cell.find('{%s}cellAddr' % ns), cell.find('{%s}cellSpan' % ns)
                            c = int(addr.get('colAddr', column)) if addr is not None else column
                            rr = int(addr.get('rowAddr', r)) if addr is not None else r
                            width = int(span.get('colSpan', 1)) if span is not None else 1
                            height = int(span.get('rowSpan', 1)) if span is not None else 1
                        else:
                            span = cell.find('.//{%s}gridSpan' % ns)
                            c, rr, height = column, r, 1
                            width = int(span.get('{%s}val' % ns, 1)) if span is not None else 1
                        layout.append((cell, rr, c, height, width))
                        column = c + width
                for row_no, row in enumerate(rows):
                    cells = row.findall('{%s}tc' % ns)
                    row_context = ' | '.join(node_text(cell, ns)[:300] or '[빈칸]' for cell in cells)
                    for col, cell in enumerate(cells):
                        if any(local_name(n) in ('tbl', 'pic', 'drawing', 'object') for n in cell.iterdescendants()):
                            continue
                        tid = f'{path}:t{table_no}r{row_no}c{col}'
                        previous = rows[row_no-1] if row_no else None
                        above = ' | '.join(node_text(c, ns)[:100] for c in previous.findall('{%s}tc' % ns)) if previous is not None else ''
                        text = node_text(cell, ns)
                        _, rr, cc, _, _ = next(item for item in layout if item[0] is cell)
                        left = sorted((item for item in layout if item[1] <= rr < item[1] + item[3] and item[2] + item[4] <= cc), key=lambda item: item[2])
                        labels = [node_text(item[0], ns)[:200] for item in left if node_text(item[0], ns).strip() and not editable_text(node_text(item[0], ns))]
                        size = cell.find('{%s}cellSz' % ns)
                        spacer = not text and not labels and size is not None and int(size.get('height', '9999')) < 1500
                        targets.append({'id': tid, 'kind': 'cell', 'text': text[:5000], 'context': row_context[:1800], 'above': above[:600],
                                        'left_labels': labels[-4:], 'input_label': labels[-1] if labels else '', 'column': cc, 'row': rr,
                                        'editable': editable_text(text) and not spacer})
                        nodes[tid] = (cell, path)
                table_no += 1
            for index, paragraph in enumerate(root.findall('.//{%s}p' % ns)):
                if any(local_name(parent) in ('tc', 'tbl') for parent in paragraph.iterancestors()):
                    continue
                if any(local_name(child) in ('tbl', 'pic', 'ctrl', 'drawing', 'secPr') for child in paragraph.iterdescendants()):
                    continue
                text = node_text(paragraph, ns)
                tid = f'{path}:p{index}'
                targets.append({'id': tid, 'kind': 'paragraph', 'text': text[:5000], 'context': '', 'above': '', 'editable': bool(re.search(r'_{2,}|\{\{', text)) and editable_text(text)})
                nodes[tid] = (paragraph, path)
    return targets, nodes, roots


def inspect_document(name, raw):
    suffix = validate_file(name, raw)
    converted = None
    if suffix == '.hwp':
        converted = raw = convert_hwp(raw); suffix = '.hwpx'
    if suffix in ('.hwpx', '.docx'):
        targets, _, roots = xml_targets(raw, suffix)
        ns = HP if suffix == '.hwpx' else W
        text = '\n'.join(node_text(r, ns) for r in roots.values())
    elif suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw))
        targets, lines = [], []
        for sheet in book:
            if sheet.max_row > 2000 or sheet.max_column > 80:
                raise ValueError('신청 양식은 시트당 2,000행·80열 이하로 등록해주세요.')
            for row in sheet:
                context = ' | '.join(f'{c.coordinate}: {c.value}' for c in row if c.value is not None)
                if context: lines.append(sheet.title + ' ' + context)
                for cell in row:
                    if cell.data_type == 'f' or isinstance(cell, openpyxl.cell.cell.MergedCell): continue
                    if context or cell.value is not None:
                        targets.append({'id': f'{sheet.title}!{cell.coordinate}', 'kind': 'cell', 'text': str(cell.value or ''), 'context': context[:1800], 'above': '', 'editable': editable_text(str(cell.value or ''))})
        text = '\n'.join(lines)
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 150: raise ValueError('PDF는 150쪽 이하로 등록해주세요.')
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        targets = [{'id': key, 'kind': 'pdf-field', 'text': str(value.get('/V', '')), 'context': str(value.get('/TU', key)), 'above': '', 'editable': True}
                   for key, value in (reader.get_fields() or {}).items() if value.get('/FT') == '/Tx']
    elif suffix in ('.txt', '.md'):
        text = raw.decode('utf-8-sig'); targets = []
    else:
        text = ''; targets = []
    if len(targets) > 1200: raise ValueError('문서 항목이 너무 많습니다. 신청 양식만 나눠 등록해주세요.')
    return {'format': suffix[1:], 'text': text[:100000], 'targets': targets, 'converted': converted,
            'mode': 'template' if targets else 'outline'}


def replace_node(node, value, ns):
    paragraphs = [node] if local_name(node) == 'p' else node.findall('.//{%s}p' % ns)
    if not paragraphs: raise ValueError('입력할 문단을 찾지 못했습니다.')
    first = paragraphs[0]
    run_tag = 'run' if ns == HP else 'r'
    template_run = first.find('{%s}%s' % (ns, run_tag))
    # Preserve paragraph/character formatting; invalidate stale HWPX line-layout caches.
    for paragraph in paragraphs:
        for child in list(paragraph):
            if local_name(child) in (run_tag, 'linesegarray'):
                paragraph.remove(child)
    run = etree.SubElement(first, '{%s}%s' % (ns, run_tag), dict(template_run.attrib) if template_run is not None else {})
    if ns == W and template_run is not None:
        properties = template_run.find('{%s}rPr' % W)
        if properties is not None: run.append(copy.deepcopy(properties))
    if ns == W:
        properties = run.find('{%s}rPr' % W)
        if properties is None: properties = etree.SubElement(run, '{%s}rPr' % W)
        color = properties.find('{%s}color' % W)
        if color is None: color = etree.SubElement(properties, '{%s}color' % W)
        color.set('{%s}val' % W, '000000')
    if ns == HP and 'charPrIDRef' not in run.attrib: run.set('charPrIDRef', '0')
    if ns == HP:
        text = etree.SubElement(run, '{%s}t' % ns)
        lines = value.split('\n'); text.text = lines[0]
        for line in lines[1:]: etree.SubElement(text, '{%s}lineBreak' % ns).tail = line
    else:
        for i, line in enumerate(value.split('\n')):
            if i: etree.SubElement(run, '{%s}br' % ns)
            text = etree.SubElement(run, '{%s}t' % ns); text.text = line
            text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def fill_document(name, raw, fields):
    suffix = Path(name).suffix.lower()
    if suffix == '.hwp': raw, suffix = convert_hwp(raw), '.hwpx'
    values = {f['target_id']: f['value'] for f in fields if f.get('target_id') and f.get('value')}
    if suffix in ('.hwpx', '.docx'):
        targets, nodes, roots = xml_targets(raw, suffix)
        if any(key not in nodes for key in values): raise ValueError('양식 버전이 변경되었습니다. 새 초안을 생성해주세요.')
        allowed = {t['id'] for t in targets if t['editable']}
        if any(key not in allowed for key in values): raise ValueError('양식 제목·항목명·동의란은 자동 입력으로 변경할 수 없습니다.')
        ns = HP if suffix == '.hwpx' else W
        for key, value in values.items(): replace_node(nodes[key][0], value, ns)
        if suffix == '.hwpx' and values:
            with safe_zip(raw) as original: header = xml_root(original.read('Contents/header.xml'))
            properties = next((e for e in header.iter() if local_name(e) == 'charProperties'), None)
            if properties is not None:
                styles = {e.get('id'): e for e in properties if local_name(e) == 'charPr'}
                next_id = max((int(k) for k in styles), default=-1) + 1; black = {}
                for key in values:
                    for run in nodes[key][0].findall('.//{%s}run' % HP):
                        old_id = run.get('charPrIDRef')
                        if old_id not in styles: continue
                        if old_id not in black:
                            clone = copy.deepcopy(styles[old_id]); clone.set('id', str(next_id)); clone.set('textColor', '#000000')
                            properties.append(clone); black[old_id] = str(next_id); next_id += 1
                        run.set('charPrIDRef', black[old_id])
                properties.set('itemCnt', str(len(properties)))
                roots['Contents/header.xml'] = header
        output = io.BytesIO()
        with safe_zip(raw) as original, zipfile.ZipFile(output, 'w') as dest:
            for member in original.infolist():
                if member.filename.startswith('Preview/') and suffix == '.hwpx':
                    if member.filename.endswith('.txt'):
                        dest.writestr(member, '\n'.join(node_text(r, ns) for r in roots.values()).encode('utf-8'))
                    continue
                content = etree.tostring(roots[member.filename], xml_declaration=True, encoding='UTF-8', standalone=True) if member.filename in roots else original.read(member)
                dest.writestr(member, content)
        return suffix, output.getvalue()
    if suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw))
        for key, value in values.items():
            sheet, coord = key.rsplit('!', 1)
            cell = book[sheet][coord]
            if cell.data_type == 'f': raise ValueError('예산 수식 셀은 수정할 수 없습니다.')
            cell.value = value
            if value.startswith('='): cell.data_type = 's'
        out = io.BytesIO(); book.save(out); return suffix, out.getvalue()
    if suffix == '.pdf' and values:
        from pypdf import PdfReader, PdfWriter
        reader = PdfReader(io.BytesIO(raw)); writer = PdfWriter(); writer.clone_document_from_reader(reader)
        writer.update_page_form_field_values(writer.pages, values, auto_regenerate=True)
        out = io.BytesIO(); writer.write(out); return suffix, out.getvalue()
    return '.docx', outline_docx(name, fields)


def outline_docx(title, fields):
    from docx import Document
    from docx.shared import Pt
    doc = Document(); doc.styles['Normal'].font.name = '맑은 고딕'; doc.styles['Normal'].font.size = Pt(10)
    doc.add_heading(title + ' — 항목별 초안', 0)
    doc.add_paragraph('공식 양식에 자동 입력되지 않은 항목별 작성문입니다. 원본 양식과 대조해 옮겨 작성해주세요.')
    for field in fields:
        doc.add_heading(field['label'], 2); doc.add_paragraph(field['value'] or '[보완 필요]')
    out = io.BytesIO(); doc.save(out); return out.getvalue()
